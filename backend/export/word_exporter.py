"""
Word文档导出模块
支持A4和A3试卷排版
"""
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os


class WordExporter:
    """Word文档导出器"""

    # 纸张尺寸配置（单位：厘米）
    PAPER_SIZES = {
        'A4': {
            'width': Cm(21),
            'height': Cm(29.7),
            'margin_top': Cm(2.54),
            'margin_bottom': Cm(2.54),
            'margin_left': Cm(3.18),
            'margin_right': Cm(3.18)
        },
        'A3': {
            'width': Cm(29.7),
            'height': Cm(42),
            'margin_top': Cm(2.5),
            'margin_bottom': Cm(2.5),
            'margin_left': Cm(2.5),
            'margin_right': Cm(2.5)
        }
    }

    def __init__(self):
        self.doc = None

    def create_exam_document(self, exam_data: Dict[str, Any],
                            output_path: str,
                            show_answers: bool = False,
                            show_analysis: bool = False) -> str:
        """
        创建试卷文档
        :param exam_data: 试卷数据
        :param output_path: 输出路径
        :param show_answers: 是否显示答案
        :param show_analysis: 是否显示解析
        :return: 文件路径
        """
        # 创建文档
        self.doc = Document()

        # 设置纸张大小
        paper_size = exam_data.get('paper_size', 'A4')
        self._set_paper_size(paper_size)

        # 设置中文字体
        self._set_chinese_font()

        # 添加试卷标题
        self._add_title(exam_data.get('name', '试卷'))

        # 添加试卷信息
        self._add_exam_info(exam_data)

        # 添加分隔线
        self._add_separator()

        # 添加题目
        questions = exam_data.get('questions', [])
        self._add_questions(questions, show_answers, show_analysis)

        # 如果显示答案，在末尾添加答案部分
        if show_answers:
            self._add_answer_section(questions)

        # 保存文档
        self.doc.save(output_path)
        return output_path

    def _set_paper_size(self, paper_size: str):
        """设置纸张大小"""
        if paper_size not in self.PAPER_SIZES:
            paper_size = 'A4'

        size_config = self.PAPER_SIZES[paper_size]

        section = self.doc.sections[0]
        section.page_width = size_config['width']
        section.page_height = size_config['height']
        section.top_margin = size_config['margin_top']
        section.bottom_margin = size_config['margin_bottom']
        section.left_margin = size_config['margin_left']
        section.right_margin = size_config['margin_right']

    def _set_chinese_font(self):
        """设置中文字体"""
        self.doc.styles['Normal'].font.name = '宋体'
        self.doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    def _add_title(self, title: str):
        """添加标题"""
        heading = self.doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 设置标题字体
        run = heading.runs[0]
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    def _add_exam_info(self, exam_data: Dict[str, Any]):
        """添加试卷信息"""
        info_paragraph = self.doc.add_paragraph()

        # 添加基本信息
        total_score = exam_data.get('total_score', 100)
        question_count = exam_data.get('question_count', 0)

        info_text = f"总分：{total_score}分    题目数量：{question_count}题    "
        info_text += "考试时间：______分钟    姓名：__________    班级：__________"

        run = info_paragraph.add_run(info_text)
        run.font.size = Pt(10.5)

        # 如果有试卷说明，添加说明
        description = exam_data.get('description')
        if description:
            desc_para = self.doc.add_paragraph()
            desc_run = desc_para.add_run(f"说明：{description}")
            desc_run.font.size = Pt(10.5)

    def _add_separator(self):
        """添加分隔线"""
        self.doc.add_paragraph('_' * 80)

    def _add_questions(self, questions: list, show_answers: bool, show_analysis: bool):
        """添加题目"""
        # 按题型分组
        questions_by_type = {}
        for q in questions:
            q_type = q.get('question_type', '其他')
            if q_type not in questions_by_type:
                questions_by_type[q_type] = []
            questions_by_type[q_type].append(q)

        # 按题型添加题目
        for q_type, type_questions in questions_by_type.items():
            # 添加题型标题
            type_heading = self.doc.add_heading(f"{q_type}", level=2)
            type_heading.runs[0].font.size = Pt(14)
            type_heading.runs[0].font.name = '黑体'
            type_heading.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

            # 添加该题型的题目
            for idx, question in enumerate(type_questions, 1):
                self._add_single_question(
                    question,
                    idx,
                    show_answers,
                    show_analysis
                )

            # 添加空行
            self.doc.add_paragraph()

    def _add_single_question(self, question: Dict[str, Any], index: int,
                            show_answers: bool, show_analysis: bool):
        """添加单个题目"""
        # 题目内容
        question_para = self.doc.add_paragraph()
        question_para.paragraph_format.left_indent = Cm(0.5)

        # 题号和分数
        score = question.get('score', 0)
        question_run = question_para.add_run(f"{index}. ")
        question_run.font.bold = True
        question_run.font.size = Pt(10.5)

        # 题目内容
        content = question.get('content', '')
        content_run = question_para.add_run(content)
        content_run.font.size = Pt(10.5)

        # 分数标注
        if score:
            score_run = question_para.add_run(f"  ({score}分)")
            score_run.font.size = Pt(9)
            score_run.font.color.rgb = RGBColor(128, 128, 128)

        # 答题区域（如果不显示答案）
        if not show_answers:
            answer_space = self.doc.add_paragraph()
            answer_space.paragraph_format.left_indent = Cm(1)
            answer_space.add_run("答：_" + "_" * 60)

        # 答案（如果需要显示）
        if show_answers:
            answer = question.get('answer', '')
            if answer:
                answer_para = self.doc.add_paragraph()
                answer_para.paragraph_format.left_indent = Cm(1)
                answer_label = answer_para.add_run("答案：")
                answer_label.font.bold = True
                answer_label.font.size = Pt(10.5)
                answer_label.font.color.rgb = RGBColor(255, 0, 0)

                answer_content = answer_para.add_run(answer)
                answer_content.font.size = Pt(10.5)

        # 解析（如果需要显示）
        if show_analysis:
            analysis = question.get('analysis', '')
            if analysis:
                analysis_para = self.doc.add_paragraph()
                analysis_para.paragraph_format.left_indent = Cm(1)
                analysis_label = analysis_para.add_run("解析：")
                analysis_label.font.bold = True
                analysis_label.font.size = Pt(10.5)
                analysis_label.font.color.rgb = RGBColor(0, 0, 255)

                analysis_content = analysis_para.add_run(analysis)
                analysis_content.font.size = Pt(10.5)

    def _add_answer_section(self, questions: list):
        """添加答案部分"""
        # 添加分页符
        self.doc.add_page_break()

        # 添加答案标题
        answer_heading = self.doc.add_heading("参考答案", level=1)
        answer_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        answer_heading.runs[0].font.size = Pt(18)

        # 按题型分组
        questions_by_type = {}
        for q in questions:
            q_type = q.get('question_type', '其他')
            if q_type not in questions_by_type:
                questions_by_type[q_type] = []
            questions_by_type[q_type].append(q)

        # 添加各题型答案
        for q_type, type_questions in questions_by_type.items():
            type_heading = self.doc.add_heading(f"{q_type}", level=2)
            type_heading.runs[0].font.size = Pt(12)

            for idx, question in enumerate(type_questions, 1):
                answer = question.get('answer', '')
                if answer:
                    answer_para = self.doc.add_paragraph()
                    answer_para.add_run(f"{idx}. ").bold = True
                    answer_para.add_run(answer)

    def export_question_bank(self, questions: list, output_path: str,
                            title: str = "题库") -> str:
        """
        导出题库到Word文档
        :param questions: 题目列表
        :param output_path: 输出路径
        :param title: 文档标题
        :return: 文件路径
        """
        self.doc = Document()
        self._set_paper_size('A4')
        self._set_chinese_font()

        # 添加标题
        self._add_title(title)

        # 添加题目
        for idx, question in enumerate(questions, 1):
            # 题目
            q_para = self.doc.add_paragraph()
            q_para.add_run(f"{idx}. ").bold = True
            q_para.add_run(question.get('content', ''))

            # 标签
            tags = question.get('tags', [])
            if tags:
                tag_para = self.doc.add_paragraph()
                tag_para.paragraph_format.left_indent = Cm(0.5)
                tag_run = tag_para.add_run(f"标签：{', '.join(tags)}")
                tag_run.font.size = Pt(9)
                tag_run.font.color.rgb = RGBColor(100, 100, 100)

            # 答案
            answer = question.get('answer')
            if answer:
                ans_para = self.doc.add_paragraph()
                ans_para.paragraph_format.left_indent = Cm(0.5)
                ans_para.add_run("答案：").bold = True
                ans_para.add_run(answer)

            # 分隔
            self.doc.add_paragraph()

        # 保存
        self.doc.save(output_path)
        return output_path
